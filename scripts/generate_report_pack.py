"""One command regenerates the pivot-ready Excel report pack: the 32-rule
catalog, a flat per-rule result row, and a flat per-violation detail row,
each sheet shaped so a business partner can drop it straight into an
Excel PivotTable without reshaping anything first.

    python scripts/generate_report_pack.py [--seed-violations]

By default this runs the rule catalog against a fresh clean load (no
seeded violations); pass --seed-violations to run it against the same
32-seeded-violation load the gate benchmark measures, so the pack's
"Violations Detail" sheet has 32 real rows to show, not zero.
"""

from __future__ import annotations

import sys
from datetime import datetime
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from openpyxl import Workbook

from app import db, engine, generator, violations
from app.rules import RULES

OUT_PATH = Path(__file__).resolve().parents[1] / "docs" / "report_pack.xlsx"


def build_pack(conn, run_label: str) -> Workbook:
    found = engine.run_all(conn)
    violation_counts: dict[str, int] = {}
    for v in found:
        violation_counts[v.rule_id] = violation_counts.get(v.rule_id, 0) + 1

    wb = Workbook()

    catalog = wb.active
    catalog.title = "Rule Catalog"
    catalog.append(["rule_id", "rule_name", "table", "owner", "success_criterion"])
    for r in RULES:
        catalog.append([r.rule_id, r.name, r.table, r.owner, r.success_criterion])

    results = wb.create_sheet("Rule Results")
    results.append(["run_label", "rule_id", "rule_name", "table", "owner", "status", "violation_count"])
    for r in RULES:
        count = violation_counts.get(r.rule_id, 0)
        results.append([run_label, r.rule_id, r.name, r.table, r.owner,
                         "FAIL" if count else "PASS", count])

    detail = wb.create_sheet("Violations Detail")
    detail.append(["run_label", "rule_id", "rule_name", "table", "owner", "record_id"])
    for v in found:
        detail.append([run_label, v.rule_id, v.rule_name, v.table, v.owner, v.record_id])

    return wb


def main() -> int:
    seed_violations = "--seed-violations" in sys.argv

    conn = db.connect()
    if seed_violations:
        clean = generator.generate_clean_load(seed=1)
        mutated, _ = violations.inject_seeded_violations(clean, seed=777)
        db.load(conn, mutated)
        run_label = "seeded-violation demo load"
    else:
        data = generator.generate_clean_load(seed=1)
        db.load(conn, data)
        run_label = "clean load"

    wb = build_pack(conn, run_label)
    OUT_PATH.parent.mkdir(exist_ok=True)
    wb.save(OUT_PATH)
    print(f"wrote {OUT_PATH} ({run_label}) at {datetime.now().isoformat(timespec='seconds')}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
